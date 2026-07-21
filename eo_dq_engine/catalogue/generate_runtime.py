"""
Enterprise Observability Platform

Component:
    EO-DQ Engine

Module:
    generate_runtime.py

Description:
    Generates runtime configuration files from the
    Enterprise DQ Catalogue.

Outputs
-------
dq_rules.yaml
dq_rules.json
catalogue_summary.json

Author
------
Enterprise Observability Team

Version
-------
1.0.0-alpha
"""

from __future__ import annotations

import argparse
import json
import logging
from collections import Counter
from datetime import datetime, UTC
from pathlib import Path
from typing import Any

import openpyxl
import yaml


logging.basicConfig(
    level=logging.INFO,
    format="%(levelname)s %(message)s"
)

LOGGER = logging.getLogger(__name__)


class RuntimeGenerator:
    """
    Generates runtime configuration from the
    Enterprise DQ Catalogue.
    """

    # ---------------------------------------------------------

    def __init__(
        self,
        catalogue: str | Path,
        output: str | Path,
    ) -> None:

        self.catalogue = Path(catalogue)
        self.output = Path(output)

        self.workbook = None
        self.sheet = None

        self.rules: list[dict[str, Any]] = []

    # ---------------------------------------------------------

    def load_workbook(self) -> None:

        LOGGER.info(
            "Loading catalogue: %s",
            self.catalogue,
        )

        if not self.catalogue.exists():
            raise FileNotFoundError(self.catalogue)

        self.workbook = openpyxl.load_workbook(
            self.catalogue,
            data_only=True,
        )

        self.sheet = self.workbook["Rules"]

    # ---------------------------------------------------------

    @staticmethod
    def normalize(value: Any) -> Any:

        if value is None:
            return None

        if isinstance(value, str):

            value = value.strip()

            if value == "":
                return None

        return value

    # ---------------------------------------------------------

    @staticmethod
    def parse_list(value: Any) -> list[str]:

        value = RuntimeGenerator.normalize(value)

        if value is None:
            return []

        if isinstance(value, list):
            return value

        return [
            item.strip()
            for item in str(value).split(",")
            if item.strip()
        ]

    # ---------------------------------------------------------

    @staticmethod
    def parse_bool(value: Any) -> bool:

        if isinstance(value, bool):
            return value

        if value is None:
            return False

        return str(value).strip().upper() == "TRUE"

    # ---------------------------------------------------------

    def build_rules(self) -> None:

        LOGGER.info("Building runtime rules...")

        headers = []

        for cell in self.sheet[1]:
            headers.append(str(cell.value).strip())

        for row in self.sheet.iter_rows(min_row=2, values_only=True):

            values = dict(zip(headers, row))

            rule = {

                "rule_id":
                    self.normalize(values.get("Rule ID")),

                "name":
                    self.normalize(values.get("Name")),

                "description":
                    self.normalize(values.get("Description")),

                "category":
                    self.normalize(values.get("Category")),

                "tier":
                    self.normalize(values.get("Tier")),

                "signal":
                    self.parse_list(
                        values.get("Signal")
                    ),

                "validator":
                    str(
                        values.get("Validator")
                    ).strip().lower(),

                "enabled":
                    self.parse_bool(
                        values.get("Enabled")
                    ),

                "severity":
                    str(
                        values.get("Severity")
                    ).upper(),

                "execution_order":
                    values.get("Execution Order"),

                "field":
                    self.normalize(
                        values.get("Field")
                    ),

                "expected":
                    self.normalize(
                        values.get("Expected")
                    ),

                "remediation":
                    self.normalize(
                        values.get("Remediation")
                    ),

                "owner":
                    self.normalize(
                        values.get("Owner")
                    ),

                "version":
                    self.normalize(
                        values.get("Version")
                    ),

                "status":
                    self.normalize(
                        values.get("Status")
                    ),

                "notes":
                    self.normalize(
                        values.get("Notes")
                    ),
            }

            self.rules.append(rule)

        LOGGER.info(
            "Loaded %d runtime rules.",
            len(self.rules),
        )

        # ---------------------------------------------------------

    def build_runtime(self) -> dict[str, Any]:
        """
        Build the runtime document that will be written
        to YAML and JSON.
        """

        return {
            "version": "1.0.0",
            "catalogue_version": "1.0",
            "generated_at": datetime.now(UTC).isoformat(),
            "generator": "Enterprise Observability Runtime Generator",
            "rule_count": len(self.rules),
            "rules": self.rules,
        }

    # ---------------------------------------------------------

    def write_yaml(self) -> Path:
        """
        Write runtime rules to YAML.
        """

        self.output.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_file = self.output / "dq_rules.yaml"

        LOGGER.info(
            "Writing %s",
            output_file,
        )

        runtime = self.build_runtime()

        with open(
            output_file,
            "w",
            encoding="utf-8",
        ) as fp:

            yaml.safe_dump(
                runtime,
                fp,
                sort_keys=False,
                allow_unicode=True,
                default_flow_style=False,
            )

        return output_file

    # ---------------------------------------------------------

    def write_json(self) -> Path:
        """
        Write runtime rules to JSON.
        """

        self.output.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_file = self.output / "dq_rules.json"

        LOGGER.info(
            "Writing %s",
            output_file,
        )

        runtime = self.build_runtime()

        with open(
            output_file,
            "w",
            encoding="utf-8",
        ) as fp:

            json.dump(
                runtime,
                fp,
                indent=4,
                ensure_ascii=False,
            )

        return output_file

    # ---------------------------------------------------------

    def validator_statistics(self) -> dict[str, int]:
        """
        Count validators.
        """

        counter = Counter()

        for rule in self.rules:

            validator = rule.get("validator")

            if validator:
                counter[validator] += 1

        return dict(counter)

    # ---------------------------------------------------------

    def signal_statistics(self) -> dict[str, int]:
        """
        Count signals.
        """

        counter = Counter()

        for rule in self.rules:

            for signal in rule.get("signal", []):

                counter[signal] += 1

        return dict(counter)
    
        # ---------------------------------------------------------

    def build_summary(self) -> dict[str, Any]:
        """
        Build runtime catalogue summary.
        """

        enabled = sum(
            1 for rule in self.rules
            if rule["enabled"]
        )

        disabled = len(self.rules) - enabled

        category_counter = Counter()

        for rule in self.rules:

            category = rule.get("category")

            if category:
                category_counter[category] += 1

        return {

            "catalogue_version": "1.0",

            "generated_at":
                datetime.now(UTC).isoformat(),

            "total_rules":
                len(self.rules),

            "enabled_rules":
                enabled,

            "disabled_rules":
                disabled,

            "validators":
                self.validator_statistics(),

            "signals":
                self.signal_statistics(),

            "categories":
                dict(category_counter),
        }

    # ---------------------------------------------------------

    def write_summary(self) -> Path:
        """
        Write catalogue summary.
        """

        self.output.mkdir(
            parents=True,
            exist_ok=True,
        )

        output_file = (
            self.output /
            "catalogue_summary.json"
        )

        LOGGER.info(
            "Writing %s",
            output_file,
        )

        summary = self.build_summary()

        with open(
            output_file,
            "w",
            encoding="utf-8",
        ) as fp:

            json.dump(
                summary,
                fp,
                indent=4,
            )

        return output_file

    # ---------------------------------------------------------

    def generate(self) -> None:
        """
        Generate all runtime artifacts.
        """

        LOGGER.info(
            "Generating runtime configuration..."
        )

        self.load_workbook()

        self.build_rules()

        self.write_yaml()

        self.write_json()

        self.write_summary()

        LOGGER.info("------------------------------------------")
        LOGGER.info("Runtime Generation Summary")
        LOGGER.info("------------------------------------------")
        LOGGER.info("Rules      : %d", len(self.rules))
        LOGGER.info("YAML       : %s", self.output / "dq_rules.yaml")
        LOGGER.info("JSON       : %s", self.output / "dq_rules.json")
        LOGGER.info("Summary    : %s", self.output / "catalogue_summary.json")
        LOGGER.info("------------------------------------------")

# =============================================================
# CLI
# =============================================================

def parse_args():

    parser = argparse.ArgumentParser(
        description="Generate EO Runtime Configuration"
    )

    parser.add_argument(
        "--catalogue",
        required=True,
        help="Enterprise_DQ_Rules.xlsx",
    )

    parser.add_argument(
        "--output",
        required=True,
        help="Output directory",
    )

    return parser.parse_args()


# =============================================================

def main():

    args = parse_args()

    generator = RuntimeGenerator(
        catalogue=args.catalogue,
        output=args.output,
    )

    generator.generate()


# =============================================================

if __name__ == "__main__":

    main()